from collections import Counter
from datetime import date, datetime, timedelta
from io import BytesIO
from math import ceil
from statistics import median

from django.db.models import Count, Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import (
    AuditCase,
    DeadlineExtension,
    Evidence,
    Finding,
    Recommendation,
    Response,
)
from .services import with_effective_deadline
from .statistics import RESPONSE_DUE_STATUSES, TERMINAL_RECOMMENDATION_STATUSES


NAVY = "102A43"
BLUE = "145DA0"
LIGHT_BLUE = "EAF3FB"
LIGHT_GRAY = "F4F7F9"
MID_GRAY = "5F6B76"
LINE = "D7E0E8"
WHITE = "FFFFFF"
GREEN = "2E7D32"
LIGHT_GREEN = "E8F5E9"
AMBER = "A85D00"
LIGHT_AMBER = "FFF3E0"
RED = "B42318"
LIGHT_RED = "FDECEC"

THIN_LINE = Side(style="thin", color=LINE)
SECTION_BORDER = Border(bottom=Side(style="medium", color=BLUE))

MONTH_NAMES = (
    "ene",
    "feb",
    "mar",
    "abr",
    "may",
    "jun",
    "jul",
    "ago",
    "sep",
    "oct",
    "nov",
    "dic",
)


def _display_user(user):
    if not user:
        return ""
    return user.get_full_name() or user.username


def _safe_text(value, max_length=32767):
    """Keep database text as text in Excel and avoid formula injection."""
    if value is None:
        return ""
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if text.lstrip().startswith(("=", "+", "-", "@")):
        text = f"'{text}"
    if len(text) > max_length:
        text = f"{text[: max_length - 24]} … [texto truncado]"
    return text


def _excel_datetime(value):
    if not value:
        return None
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.replace(tzinfo=None)


def _month_label(value):
    if not value:
        return ""
    return f"{MONTH_NAMES[value.month - 1]} {value.year}"


def _format_filter_value(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    return value or "Todos"


def _write_section_title(worksheet, row, title, end_column):
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=end_column,
    )
    cell = worksheet.cell(row=row, column=1, value=title)
    cell.font = Font(name="Aptos Display", size=14, bold=True, color=NAVY)
    cell.border = SECTION_BORDER
    cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[row].height = 24


def _style_table_header(row_cells):
    for cell in row_cells:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=BLUE))


def _add_table(worksheet, reference, name):
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _apply_sheet_setup(worksheet, landscape=True):
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_setup.orientation = "landscape" if landscape else "portrait"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LETTER
    worksheet.print_title_rows = "1:4"
    worksheet.sheet_properties.outlinePr.summaryBelow = True
    worksheet.oddFooter.center.text = "Página &P de &N"
    worksheet.oddFooter.right.text = "Dirección de Auditoría Interna"
    worksheet.oddFooter.center.size = 8
    worksheet.oddFooter.right.size = 8


def _create_data_sheet(
    workbook,
    *,
    title,
    description,
    headers,
    rows,
    table_name,
    widths=None,
    date_columns=(),
    datetime_columns=(),
    percentage_columns=(),
    integer_columns=(),
    wrap_columns=(),
    landscape=True,
):
    worksheet = workbook.create_sheet(title=title)
    end_column = len(headers)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    title_cell = worksheet.cell(row=1, column=1, value=title)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    title_cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 34

    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    description_cell = worksheet.cell(row=2, column=1, value=description)
    description_cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    description_cell.font = Font(name="Aptos", size=9, color=NAVY, italic=True)
    description_cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[2].height = 30

    header_row = 4
    for column, header in enumerate(headers, start=1):
        worksheet.cell(row=header_row, column=column, value=header)
    _style_table_header(worksheet[header_row])
    worksheet.row_dimensions[header_row].height = 34

    for row_index, values in enumerate(rows, start=header_row + 1):
        for column_index, value in enumerate(values, start=1):
            if isinstance(value, str):
                value = _safe_text(value)
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            cell.font = Font(name="Aptos", size=9, color="17212B")
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column_index in wrap_columns,
            )
            if column_index in date_columns and value:
                cell.number_format = "dd/mm/yyyy"
            elif column_index in datetime_columns and value:
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif column_index in percentage_columns and value is not None:
                cell.number_format = "0%"
                cell.alignment = Alignment(horizontal="right", vertical="top")
            elif column_index in integer_columns and value is not None:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="top")

    last_row = header_row + len(rows)
    if rows:
        reference = f"A{header_row}:{get_column_letter(end_column)}{last_row}"
        _add_table(worksheet, reference, table_name)
    else:
        worksheet.merge_cells(
            start_row=header_row + 1,
            start_column=1,
            end_row=header_row + 1,
            end_column=end_column,
        )
        empty_cell = worksheet.cell(row=header_row + 1, column=1, value="Sin registros para los filtros seleccionados.")
        empty_cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        empty_cell.font = Font(name="Aptos", size=10, italic=True, color=MID_GRAY)
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[header_row + 1].height = 28

    widths = widths or {}
    for column_index, header in enumerate(headers, start=1):
        configured = widths.get(column_index)
        if configured:
            width = configured
        else:
            lengths = [len(str(header))]
            for values in rows[:100]:
                value = values[column_index - 1]
                if value is not None:
                    lengths.append(len(str(value).split("\n", 1)[0]))
            width = min(max(max(lengths) + 2, 11), 35)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    _apply_sheet_setup(worksheet, landscape=landscape)
    worksheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(end_column)}{max(last_row, header_row)}"
    )
    worksheet.print_area = f"A1:{get_column_letter(end_column)}{max(last_row, header_row + 1)}"
    return worksheet


def _build_dataset(cases, today):
    case_list = list(
        cases.select_related("audited_organization", "assigned_auditor")
        .order_by("created_at", "pk")
    )
    case_ids = [case.pk for case in case_list]

    findings = list(
        Finding.objects.filter(case_id__in=case_ids)
        .select_related("case__audited_organization", "case__assigned_auditor")
        .annotate(recommendation_count=Count("recommendations"))
        .order_by("case__reference", "number", "pk")
    )
    recommendations = list(
        with_effective_deadline(
            Recommendation.objects.filter(finding__case_id__in=case_ids)
        )
        .select_related(
            "finding__case__audited_organization",
            "finding__case__assigned_auditor",
            "responsible_organization",
            "source_recommendation__source_document",
            "carried_from",
        )
        .annotate(
            response_count=Count("responses", distinct=True),
            extension_count=Count("deadline_extensions", distinct=True),
        )
        .order_by("finding__case__reference", "finding__number", "number", "pk")
    )
    recommendation_ids = [item.pk for item in recommendations]
    responses = list(
        Response.objects.filter(recommendation_id__in=recommendation_ids)
        .select_related(
            "recommendation__finding__case__audited_organization",
            "recommendation__finding__case__assigned_auditor",
            "submitted_by",
            "school_board_period",
            "review__reviewed_by",
        )
        .annotate(
            evidence_count=Count("evidence", distinct=True),
            evidence_pending_count=Count(
                "evidence",
                filter=Q(evidence__scan_status=Evidence.ScanStatus.PENDING),
                distinct=True,
            ),
            evidence_clean_count=Count(
                "evidence",
                filter=Q(evidence__scan_status=Evidence.ScanStatus.CLEAN),
                distinct=True,
            ),
            evidence_rejected_count=Count(
                "evidence",
                filter=Q(evidence__scan_status=Evidence.ScanStatus.REJECTED),
                distinct=True,
            ),
        )
        .order_by("recommendation_id", "version", "pk")
    )
    extensions = list(
        DeadlineExtension.objects.filter(recommendation_id__in=recommendation_ids)
        .select_related(
            "recommendation__finding__case__audited_organization",
            "recommendation__finding__case__assigned_auditor",
            "granted_by",
        )
        .order_by("recommendation__finding__case__reference", "granted_at", "pk")
    )

    case_finding_count = Counter(item.case_id for item in findings)
    case_recommendation_count = Counter(item.finding.case_id for item in recommendations)
    case_response_count = Counter(
        item.recommendation.finding.case_id for item in responses
    )
    case_pending_review_count = Counter(
        item.recommendation.finding.case_id
        for item in responses
        if not hasattr(item, "review")
    )
    case_high_risk_count = Counter(
        item.case_id
        for item in findings
        if item.risk_level in (Finding.RiskLevel.HIGH, Finding.RiskLevel.CRITICAL)
    )
    case_overdue_count = Counter(
        item.finding.case_id
        for item in recommendations
        if item.current_deadline
        and item.current_deadline < today
        and item.status in RESPONSE_DUE_STATUSES
    )
    case_terminal_count = Counter(
        item.finding.case_id
        for item in recommendations
        if item.status in TERMINAL_RECOMMENDATION_STATUSES
    )
    case_complied_count = Counter(
        item.finding.case_id
        for item in recommendations
        if item.status == Recommendation.Status.COMPLIED
    )

    latest_response = {}
    for response in responses:
        latest_response[response.recommendation_id] = response

    return {
        "cases": case_list,
        "findings": findings,
        "recommendations": recommendations,
        "responses": responses,
        "extensions": extensions,
        "latest_response": latest_response,
        "case_finding_count": case_finding_count,
        "case_recommendation_count": case_recommendation_count,
        "case_response_count": case_response_count,
        "case_pending_review_count": case_pending_review_count,
        "case_high_risk_count": case_high_risk_count,
        "case_overdue_count": case_overdue_count,
        "case_terminal_count": case_terminal_count,
        "case_complied_count": case_complied_count,
    }


def _case_rows(dataset):
    rows = []
    for case in dataset["cases"]:
        terminal = dataset["case_terminal_count"][case.pk]
        complied = dataset["case_complied_count"][case.pk]
        rows.append(
            [
                case.reference,
                case.title,
                case.audited_organization.code,
                case.audited_organization.name,
                case.audited_organization.department,
                case.audited_organization.municipality,
                _display_user(case.assigned_auditor),
                case.assigned_auditor.username,
                case.get_status_display(),
                case.report_date,
                case.period_start,
                case.period_end,
                case.response_deadline,
                _excel_datetime(case.created_at),
                _excel_datetime(case.updated_at),
                dataset["case_finding_count"][case.pk],
                dataset["case_high_risk_count"][case.pk],
                dataset["case_recommendation_count"][case.pk],
                dataset["case_overdue_count"][case.pk],
                dataset["case_response_count"][case.pk],
                dataset["case_pending_review_count"][case.pk],
                complied / terminal if terminal else None,
            ]
        )
    return rows


def _finding_rows(dataset):
    return [
        [
            finding.case.reference,
            finding.case.audited_organization.code,
            finding.case.audited_organization.name,
            _display_user(finding.case.assigned_auditor),
            finding.number,
            finding.title,
            finding.get_risk_level_display(),
            finding.condition,
            finding.criteria,
            finding.cause,
            finding.effect,
            finding.recommendation_count,
        ]
        for finding in dataset["findings"]
    ]


def _recommendation_rows(dataset, today):
    rows = []
    for recommendation in dataset["recommendations"]:
        response = dataset["latest_response"].get(recommendation.pk)
        review = response.review if response and hasattr(response, "review") else None
        effective_deadline = recommendation.current_deadline
        is_due_status = recommendation.status in RESPONSE_DUE_STATUSES
        overdue_days = (
            (today - effective_deadline).days
            if effective_deadline and effective_deadline < today and is_due_status
            else 0
        )
        days_to_due = (
            (effective_deadline - today).days
            if effective_deadline and effective_deadline >= today and is_due_status
            else None
        )
        source = ""
        if recommendation.source_recommendation_id:
            document = recommendation.source_recommendation.source_document
            source = f"Informe anterior: {document.reference or document.title}"
        elif recommendation.carried_from_id:
            source = f"Recomendación anterior: {recommendation.carried_from}"

        rows.append(
            [
                recommendation.finding.case.reference,
                recommendation.finding.case.audited_organization.code,
                recommendation.finding.case.audited_organization.name,
                _display_user(recommendation.finding.case.assigned_auditor),
                recommendation.finding.number,
                recommendation.finding.title,
                recommendation.finding.get_risk_level_display(),
                recommendation.number,
                recommendation.text,
                recommendation.responsible_organization.code,
                recommendation.responsible_organization.name,
                recommendation.get_status_display(),
                recommendation.deadline,
                effective_deadline,
                overdue_days,
                days_to_due,
                recommendation.response_count,
                recommendation.extension_count,
                _excel_datetime(response.submitted_at) if response else None,
                response.get_declared_status_display() if response else "",
                review.get_outcome_display() if review else "",
                _excel_datetime(review.reviewed_at) if review else None,
                _display_user(review.reviewed_by) if review else "",
                recommendation.evidence_requirements,
                source,
                _excel_datetime(recommendation.no_response_recorded_at),
            ]
        )
    return rows


def _response_rows(dataset):
    rows = []
    for response in dataset["responses"]:
        review = response.review if hasattr(response, "review") else None
        school_board = response.school_board_period
        review_hours = (
            (review.reviewed_at - response.submitted_at).total_seconds() / 3600
            if review
            else None
        )
        rows.append(
            [
                response.recommendation.finding.case.reference,
                response.recommendation.finding.case.audited_organization.code,
                response.recommendation.finding.case.audited_organization.name,
                _display_user(response.recommendation.finding.case.assigned_auditor),
                response.recommendation.finding.number,
                response.recommendation.number,
                response.version,
                response.get_declared_status_display(),
                response.action_description,
                response.action_date,
                response.responsible_name,
                response.responsible_job_title,
                response.non_compliance_reason,
                response.action_plan,
                response.expected_completion_date,
                "Sí" if response.accuracy_declaration else "No",
                _display_user(response.submitted_by),
                _excel_datetime(response.submitted_at),
                response.evidence_count,
                response.evidence_pending_count,
                response.evidence_clean_count,
                response.evidence_rejected_count,
                (
                    f"{school_board.school_year_start}-{school_board.school_year_end}"
                    if school_board
                    else ""
                ),
                review.get_outcome_display() if review else "Pendiente de revisión",
                review.comments if review else "",
                _display_user(review.reviewed_by) if review else "",
                _excel_datetime(review.reviewed_at) if review else None,
                review_hours,
                review_hours / 24 if review_hours is not None else None,
            ]
        )
    return rows


def _extension_rows(dataset):
    return [
        [
            extension.recommendation.finding.case.reference,
            extension.recommendation.finding.case.audited_organization.code,
            extension.recommendation.finding.case.audited_organization.name,
            _display_user(extension.recommendation.finding.case.assigned_auditor),
            extension.recommendation.finding.number,
            extension.recommendation.number,
            extension.previous_deadline,
            extension.business_days,
            extension.new_deadline,
            extension.reason,
            _display_user(extension.granted_by),
            _excel_datetime(extension.granted_at),
        ]
        for extension in dataset["extensions"]
    ]


def _auditor_rows(auditor_statistics):
    output = []
    for row in auditor_statistics:
        terminal = row.get("terminal_count", 0)
        output.append(
            [
                _display_user(row["auditor"]),
                row["auditor"].username,
                row["auditor"].job_title,
                "Sí" if row["auditor"].is_active else "No",
                row["cases_count"],
                row["open_cases_count"],
                row.get("in_management_count", row["open_cases_count"]),
                row["recommendation_count"],
                row.get("complied_count", 0),
                row.get("partial_count", 0),
                row.get("not_complied_count", 0),
                row["pending_reviews_count"],
                row["overdue_count"],
                row.get("complied_count", 0) / terminal if terminal else None,
            ]
        )
    return output


def _organization_rows(organization_statistics):
    output = []
    for row in organization_statistics:
        terminal = row.get("terminal_count", 0)
        output.append(
            [
                row["code"],
                row["name"],
                row.get("kind", ""),
                row.get("department", ""),
                row.get("municipality", ""),
                row["cases_count"],
                row["open_cases_count"],
                row.get("in_management_count", row["open_cases_count"]),
                row["recommendation_count"],
                row.get("complied_count", 0),
                row.get("partial_count", 0),
                row.get("not_complied_count", 0),
                row["overdue_count"],
                row["critical_count"],
                row.get("complied_count", 0) / terminal if terminal else None,
            ]
        )
    return output


def _month_start(value):
    return date(value.year, value.month, 1)


def _next_month(value):
    if value.month == 12:
        return date(value.year + 1, 1, 1)
    return date(value.year, value.month + 1, 1)


def _series_rows(dataset, filters, today):
    case_dates = [timezone.localdate(item.created_at) for item in dataset["cases"]]
    response_dates = [timezone.localdate(item.submitted_at) for item in dataset["responses"]]
    review_dates = [
        timezone.localdate(item.review.reviewed_at)
        for item in dataset["responses"]
        if hasattr(item, "review")
    ]
    event_dates = case_dates + response_dates + review_dates
    if event_dates:
        start = min(event_dates)
        end = max(event_dates + [today])
    else:
        start = filters.get("start_date") or today
        end = filters.get("end_date") or today
    if filters.get("start_date"):
        start = min(start, filters["start_date"])
    start = _month_start(start)
    end = _month_start(end)

    cases_by_month = Counter(_month_start(value) for value in case_dates)
    responses_by_month = Counter(_month_start(value) for value in response_dates)
    reviews_by_month = Counter(_month_start(value) for value in review_dates)
    rows = []
    current = start
    while current <= end:
        rows.append(
            [
                current,
                _month_label(current),
                cases_by_month[current],
                responses_by_month[current],
                reviews_by_month[current],
            ]
        )
        current = _next_month(current)
    return rows


def _extended_analysis(dataset, filters, today):
    recommendations = dataset["recommendations"]
    terminal = [
        item for item in recommendations if item.status in TERMINAL_RECOMMENDATION_STATUSES
    ]
    complied = [item for item in terminal if item.status == Recommendation.Status.COMPLIED]
    partial = [item for item in terminal if item.status == Recommendation.Status.PARTIAL]
    not_complied = [
        item for item in terminal if item.status == Recommendation.Status.NOT_COMPLIED
    ]
    due_soon = [
        item
        for item in recommendations
        if item.current_deadline
        and today <= item.current_deadline <= today + timedelta(days=7)
        and item.status in RESPONSE_DUE_STATUSES
    ]
    no_deadline = [item for item in recommendations if not item.current_deadline]
    automatic_non_compliance = [
        item for item in recommendations if item.no_response_recorded_at is not None
    ]
    review_days = sorted(
        (item.review.reviewed_at - item.submitted_at).total_seconds() / 86400
        for item in dataset["responses"]
        if hasattr(item, "review")
    )
    pending_response_versions = [
        item for item in dataset["responses"] if not hasattr(item, "review")
    ]
    pending_recommendations = {
        item.recommendation_id for item in pending_response_versions
    }
    in_management = [
        item
        for item in dataset["cases"]
        if item.status not in (AuditCase.Status.DRAFT, AuditCase.Status.CLOSED)
    ]
    return {
        "cases_in_management": len(in_management),
        "terminal_count": len(terminal),
        "complied_count": len(complied),
        "partial_count": len(partial),
        "not_complied_count": len(not_complied),
        "compliance_ratio": len(complied) / len(terminal) if terminal else None,
        "partial_ratio": len(partial) / len(terminal) if terminal else None,
        "resolution_ratio": len(terminal) / len(recommendations) if recommendations else None,
        "automatic_non_compliance_count": len(automatic_non_compliance),
        "due_soon_count": len(due_soon),
        "no_deadline_count": len(no_deadline),
        "pending_response_versions": len(pending_response_versions),
        "pending_recommendations": len(pending_recommendations),
        "median_review_days": median(review_days) if review_days else None,
        "p90_review_days": (
            review_days[max(ceil(len(review_days) * 0.9) - 1, 0)]
            if review_days
            else None
        ),
        "series_rows": _series_rows(dataset, filters, today),
    }


def _responsible_organization_rows(dataset, today):
    groups = {}
    for recommendation in dataset["recommendations"]:
        organization = recommendation.responsible_organization
        row = groups.setdefault(
            organization.pk,
            {
                "organization": organization,
                "cases": set(),
                "statuses": Counter(),
                "overdue": 0,
                "due_soon": 0,
                "without_deadline": 0,
                "automatic_non_compliance": 0,
            },
        )
        row["cases"].add(recommendation.finding.case_id)
        row["statuses"][recommendation.status] += 1
        if not recommendation.current_deadline:
            row["without_deadline"] += 1
        elif (
            recommendation.current_deadline < today
            and recommendation.status in RESPONSE_DUE_STATUSES
        ):
            row["overdue"] += 1
        elif (
            today <= recommendation.current_deadline <= today + timedelta(days=7)
            and recommendation.status in RESPONSE_DUE_STATUSES
        ):
            row["due_soon"] += 1
        if recommendation.no_response_recorded_at:
            row["automatic_non_compliance"] += 1

    output = []
    for group in groups.values():
        statuses = group["statuses"]
        terminal = sum(statuses[value] for value in TERMINAL_RECOMMENDATION_STATUSES)
        organization = group["organization"]
        output.append(
            [
                organization.code,
                organization.name,
                organization.get_kind_display(),
                organization.department,
                organization.municipality,
                len(group["cases"]),
                sum(statuses.values()),
                statuses[Recommendation.Status.PENDING],
                statuses[Recommendation.Status.SUBMITTED],
                statuses[Recommendation.Status.UNDER_REVIEW],
                statuses[Recommendation.Status.CORRECTION_REQUIRED],
                statuses[Recommendation.Status.COMPLIED],
                statuses[Recommendation.Status.PARTIAL],
                statuses[Recommendation.Status.NOT_COMPLIED],
                statuses[Recommendation.Status.COMPLIED] / terminal if terminal else None,
                group["overdue"],
                group["due_soon"],
                group["without_deadline"],
                group["automatic_non_compliance"],
            ]
        )
    return sorted(output, key=lambda row: (-row[15], -row[18], -row[6], row[1]))


def _control_rows(dataset, statistics, organization_statistics, alert_rows):
    checks = [
        (
            "Total de expedientes",
            statistics["total_cases"],
            len(dataset["cases"]),
            "El KPI debe coincidir con la hoja Expedientes.",
        ),
        (
            "Total de hallazgos",
            statistics["total_findings"],
            len(dataset["findings"]),
            "El KPI debe coincidir con la hoja Hallazgos.",
        ),
        (
            "Total de recomendaciones",
            statistics["total_recommendations"],
            len(dataset["recommendations"]),
            "El KPI debe coincidir con la hoja Recomendaciones.",
        ),
        (
            "Distribución de expedientes",
            statistics["total_cases"],
            sum(item["count"] for item in statistics["case_status_distribution"]),
            "La distribución por estado debe sumar el total de expedientes.",
        ),
        (
            "Distribución de recomendaciones",
            statistics["total_recommendations"],
            sum(item["count"] for item in statistics["recommendation_status_distribution"]),
            "La distribución por estado debe sumar el total de recomendaciones.",
        ),
        (
            "Distribución de riesgos",
            statistics["total_findings"],
            sum(item["count"] for item in statistics["risk_distribution"]),
            "La distribución por riesgo debe sumar el total de hallazgos.",
        ),
        (
            "Instituciones auditadas",
            statistics["organizations_count"],
            len(organization_statistics),
            "La hoja Instituciones debe contener todas las instituciones del corte.",
        ),
        (
            "Identificadores únicos de expedientes",
            len(dataset["cases"]),
            len({item.reference for item in dataset["cases"]}),
            "No deben existir referencias duplicadas.",
        ),
    ]
    rows = [
        [name, expected, observed, "OK" if expected == observed else "ERROR", detail]
        for name, expected, observed, detail in checks
    ]
    rows.append(
        [
            "Alertas y controles identificados",
            len(alert_rows),
            len(alert_rows),
            "OK",
            "Conteo informativo de excepciones operativas y de calidad.",
        ]
    )
    return rows


def _alert_rows(dataset, today):
    rows = []
    for case in dataset["cases"]:
        if not case.report_date:
            rows.append(
                [
                    "Calidad de datos",
                    "Media",
                    "Expediente sin fecha oficial de informe",
                    case.reference,
                    case.audited_organization.name,
                    _display_user(case.assigned_auditor),
                    None,
                    None,
                    "Complete la fecha oficial del informe.",
                ]
            )
        if not case.period_start or not case.period_end:
            rows.append(
                [
                    "Calidad de datos",
                    "Media",
                    "Período auditado incompleto",
                    case.reference,
                    case.audited_organization.name,
                    _display_user(case.assigned_auditor),
                    None,
                    None,
                    "Registre el inicio y fin del período auditado.",
                ]
            )
        if dataset["case_finding_count"][case.pk] == 0:
            rows.append(
                [
                    "Calidad de datos",
                    "Alta",
                    "Expediente sin hallazgos",
                    case.reference,
                    case.audited_organization.name,
                    _display_user(case.assigned_auditor),
                    None,
                    None,
                    "El expediente no contiene hallazgos registrados.",
                ]
            )

    for finding in dataset["findings"]:
        if finding.recommendation_count == 0:
            rows.append(
                [
                    "Calidad de datos",
                    "Alta",
                    "Hallazgo sin recomendaciones",
                    f"{finding.case.reference} / H{finding.number}",
                    finding.case.audited_organization.name,
                    _display_user(finding.case.assigned_auditor),
                    None,
                    None,
                    finding.title,
                ]
            )
        if finding.risk_level == Finding.RiskLevel.CRITICAL:
            rows.append(
                [
                    "Atención operativa",
                    "Crítica",
                    "Hallazgo de riesgo crítico",
                    f"{finding.case.reference} / H{finding.number}",
                    finding.case.audited_organization.name,
                    _display_user(finding.case.assigned_auditor),
                    None,
                    None,
                    finding.title,
                ]
            )

    for recommendation in dataset["recommendations"]:
        identifier = (
            f"{recommendation.finding.case.reference} / "
            f"H{recommendation.finding.number} / R{recommendation.number}"
        )
        if not recommendation.current_deadline:
            rows.append(
                [
                    "Calidad de datos",
                    "Media",
                    "Recomendación sin fecha límite",
                    identifier,
                    recommendation.finding.case.audited_organization.name,
                    _display_user(recommendation.finding.case.assigned_auditor),
                    None,
                    None,
                    "La recomendación no tiene fecha límite vigente.",
                ]
            )
        elif (
            recommendation.current_deadline < today
            and recommendation.status in RESPONSE_DUE_STATUSES
        ):
            days = (today - recommendation.current_deadline).days
            rows.append(
                [
                    "Atención operativa",
                    "Crítica",
                    "Recomendación vencida",
                    identifier,
                    recommendation.finding.case.audited_organization.name,
                    _display_user(recommendation.finding.case.assigned_auditor),
                    recommendation.current_deadline,
                    days,
                    recommendation.text,
                ]
            )
        elif (
            recommendation.current_deadline >= today
            and recommendation.current_deadline <= today + timedelta(days=7)
            and recommendation.status in RESPONSE_DUE_STATUSES
        ):
            days = (recommendation.current_deadline - today).days
            rows.append(
                [
                    "Atención operativa",
                    "Alta",
                    "Recomendación próxima a vencer",
                    identifier,
                    recommendation.finding.case.audited_organization.name,
                    _display_user(recommendation.finding.case.assigned_auditor),
                    recommendation.current_deadline,
                    days,
                    recommendation.text,
                ]
            )

    for response in dataset["responses"]:
        if not hasattr(response, "review"):
            recommendation = response.recommendation
            rows.append(
                [
                    "Atención operativa",
                    "Alta",
                    "Respuesta pendiente de revisión",
                    (
                        f"{recommendation.finding.case.reference} / "
                        f"H{recommendation.finding.number} / R{recommendation.number} / "
                        f"V{response.version}"
                    ),
                    recommendation.finding.case.audited_organization.name,
                    _display_user(recommendation.finding.case.assigned_auditor),
                    None,
                    (today - timezone.localdate(response.submitted_at)).days,
                    response.action_description,
                ]
            )

    priority_order = {"Crítica": 0, "Alta": 1, "Media": 2, "Baja": 3}
    return sorted(
        rows,
        key=lambda row: (
            priority_order.get(row[1], 9),
            row[2],
            row[3],
        ),
    )


def _add_status_conditional_formatting(worksheet, column_letter, first_row, last_row):
    if last_row < first_row:
        return
    target = f"{column_letter}{first_row}:{column_letter}{last_row}"
    worksheet.conditional_formatting.add(
        target,
        FormulaRule(
            formula=[f'${column_letter}{first_row}="Cumplida"'],
            fill=PatternFill("solid", fgColor=LIGHT_GREEN),
            font=Font(color=GREEN, bold=True),
        ),
    )
    worksheet.conditional_formatting.add(
        target,
        FormulaRule(
            formula=[f'${column_letter}{first_row}="No cumplida"'],
            fill=PatternFill("solid", fgColor=LIGHT_RED),
            font=Font(color=RED, bold=True),
        ),
    )
    worksheet.conditional_formatting.add(
        target,
        FormulaRule(
            formula=[f'${column_letter}{first_row}="Parcialmente cumplida"'],
            fill=PatternFill("solid", fgColor=LIGHT_AMBER),
            font=Font(color=AMBER, bold=True),
        ),
    )


def _add_summary_chart(worksheet, chart, position, width=13.5, height=7.5):
    chart.style = 10
    chart.width = width
    chart.height = height
    if chart.legend is not None:
        chart.legend.position = "r"
    worksheet.add_chart(chart, position)


def _write_kpi_card(worksheet, start_column, row, label, value, note, fill_color=LIGHT_BLUE, number_format=None):
    end_column = start_column + 2
    for current_row in range(row, row + 3):
        worksheet.merge_cells(
            start_row=current_row,
            start_column=start_column,
            end_row=current_row,
            end_column=end_column,
        )
        for column in range(start_column, end_column + 1):
            cell = worksheet.cell(row=current_row, column=column)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.border = Border(
                left=THIN_LINE if column == start_column else Side(style=None),
                right=THIN_LINE if column == end_column else Side(style=None),
                top=THIN_LINE if current_row == row else Side(style=None),
                bottom=THIN_LINE if current_row == row + 2 else Side(style=None),
            )
    label_cell = worksheet.cell(row=row, column=start_column, value=label)
    label_cell.font = Font(name="Aptos", size=9, bold=True, color=MID_GRAY)
    label_cell.alignment = Alignment(horizontal="center", vertical="center")
    value_cell = worksheet.cell(row=row + 1, column=start_column, value=value)
    value_cell.font = Font(name="Aptos Display", size=20, bold=True, color=NAVY)
    value_cell.alignment = Alignment(horizontal="center", vertical="center")
    if number_format:
        value_cell.number_format = number_format
    note_cell = worksheet.cell(row=row + 2, column=start_column, value=note)
    note_cell.font = Font(name="Aptos", size=8, color=MID_GRAY)
    note_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[row].height = 20
    worksheet.row_dimensions[row + 1].height = 30
    worksheet.row_dimensions[row + 2].height = 24


def _insights(statistics, analysis, auditor_statistics, organization_statistics):
    statements = []
    if statistics["overdue_recommendations"]:
        statements.append(
            f"Atención inmediata: {statistics['overdue_recommendations']} recomendación(es) "
            "tienen el plazo vigente vencido y permanecen pendientes de respuesta o corrección."
        )
    else:
        statements.append("No se observan recomendaciones vencidas dentro del corte seleccionado.")
    if statistics["critical_findings"]:
        statements.append(
            f"Riesgo: existen {statistics['critical_findings']} hallazgo(s) crítico(s) y "
            f"{statistics['high_risk_findings']} de riesgo alto o crítico."
        )
    if statistics["pending_reviews"]:
        statements.append(
            f"Flujo de revisión: {analysis['pending_response_versions']} versión(es) de respuesta, "
            f"correspondientes a {analysis['pending_recommendations']} recomendación(es), esperan dictamen."
        )
    if analysis["automatic_non_compliance_count"]:
        statements.append(
            f"Incumplimiento automático: {analysis['automatic_non_compliance_count']} recomendación(es) "
            "fueron registradas como no cumplidas por falta de respuesta en plazo."
        )
    if analysis["no_deadline_count"]:
        statements.append(
            f"Calidad de datos: {analysis['no_deadline_count']} recomendación(es) no poseen fecha límite vigente."
        )
    if organization_statistics:
        top = organization_statistics[0]
        statements.append(
            f"Principal foco institucional según vencimientos y riesgo: {top['name']} "
            f"({top['overdue_count']} vencida(s), {top['critical_count']} hallazgo(s) crítico(s))."
        )
    active_auditors = [row for row in auditor_statistics if row["cases_count"]]
    if active_auditors:
        top = max(
            active_auditors,
            key=lambda row: (
                row["overdue_count"],
                row["pending_reviews_count"],
                row["open_cases_count"],
            ),
        )
        statements.append(
            f"Mayor presión operativa observada: {_display_user(top['auditor'])}, con "
            f"{top['open_cases_count']} expediente(s) no cerrado(s), "
            f"{top['pending_reviews_count']} revisión(es) pendiente(s) y "
            f"{top['overdue_count']} recomendación(es) vencida(s)."
        )
    return statements[:7]


def _create_summary_sheet(
    workbook,
    *,
    report_id,
    filters,
    statistics,
    analysis,
    auditor_statistics,
    organization_statistics,
    generated_by,
    generated_at,
):
    worksheet = workbook.active
    worksheet.title = "Resumen"
    worksheet.sheet_view.showGridLines = False
    for column in range(1, 13):
        worksheet.column_dimensions[get_column_letter(column)].width = 14

    worksheet.merge_cells("A1:L2")
    title = worksheet["A1"]
    title.value = "Informe estadístico y analítico de Auditoría Interna"
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = Font(name="Aptos Display", size=22, bold=True, color=WHITE)
    title.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 28
    worksheet.row_dimensions[2].height = 24

    worksheet.merge_cells("A3:L3")
    scope_note = worksheet["A3"]
    scope_note.value = (
        "Corte actual de los expedientes registrados en el período seleccionado. "
        "Estados, vencimientos y resultados corresponden al momento de generación."
    )
    scope_note.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    scope_note.font = Font(name="Aptos", size=9, italic=True, color=NAVY)
    scope_note.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[3].height = 28

    worksheet.merge_cells("A5:F5")
    worksheet["A5"] = "Alcance del análisis"
    worksheet.merge_cells("G5:L5")
    worksheet["G5"] = "Datos de la emisión"
    for cell in (worksheet["A5"], worksheet["G5"]):
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center")

    period_label = filters.get("selected_period_label", filters["selected_period"])
    status_label = "Todos"
    if filters.get("selected_status"):
        status_label = dict(AuditCase.Status.choices)[filters["selected_status"]]
    scope_rows = [
        ("Período", period_label),
        (
            "Fechas",
            (
                f"{_format_filter_value(filters.get('start_date'))} – "
                f"{_format_filter_value(filters.get('end_date'))}"
                if filters.get("start_date") and filters.get("end_date")
                else "Todo el historial"
            ),
        ),
        ("Auditor", _display_user(filters.get("selected_auditor")) or "Todos"),
        (
            "Institución",
            filters["selected_organization"].name
            if filters.get("selected_organization")
            else "Todas",
        ),
        ("Estado del expediente", status_label),
    ]
    emission_rows = [
        ("Folio", report_id),
        ("Generado", _excel_datetime(generated_at)),
        ("Generado por", _display_user(generated_by)),
        ("Usuario", generated_by.username),
        ("Uso", "Información interna de Auditoría"),
    ]
    for offset, (left, right) in enumerate(scope_rows, start=6):
        worksheet.cell(offset, 1, left)
        worksheet.merge_cells(start_row=offset, start_column=2, end_row=offset, end_column=6)
        worksheet.cell(offset, 2, _safe_text(right))
    for offset, (left, right) in enumerate(emission_rows, start=6):
        worksheet.cell(offset, 7, left)
        worksheet.merge_cells(start_row=offset, start_column=8, end_row=offset, end_column=12)
        worksheet.cell(offset, 8, right)
    for row in range(6, 11):
        for column in (1, 7):
            cell = worksheet.cell(row, column)
            cell.font = Font(name="Aptos", size=9, bold=True, color=NAVY)
            cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        for column in list(range(2, 7)) + list(range(8, 13)):
            cell = worksheet.cell(row, column)
            cell.font = Font(name="Aptos", size=9, color="17212B")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=LINE))
    worksheet["H7"].number_format = "dd/mm/yyyy hh:mm"

    _write_kpi_card(
        worksheet,
        1,
        12,
        "EXPEDIENTES",
        statistics["total_cases"],
        (
            f"{analysis['cases_in_management']} en gestión · "
            f"{statistics['open_cases']} no cerrados · {statistics['closed_cases']} cerrados"
        ),
    )
    _write_kpi_card(
        worksheet,
        4,
        12,
        "CUMPLIMIENTO",
        analysis["compliance_ratio"] if analysis["compliance_ratio"] is not None else "N/D",
        (
            f"{statistics['complied_recommendations']} de "
            f"{statistics['terminal_recommendations']} resultados definitivos"
        ),
        fill_color=LIGHT_GREEN,
        number_format="0.0%" if analysis["compliance_ratio"] is not None else None,
    )
    worksheet["D12"].comment = Comment(
        "Recomendaciones cumplidas divididas entre recomendaciones con resultado definitivo.",
        "SIGA-MINEDUCYT",
    )
    _write_kpi_card(
        worksheet,
        7,
        12,
        "VERSIONES POR REVISAR",
        analysis["pending_response_versions"],
        f"De {analysis['pending_recommendations']} recomendación(es)",
        fill_color=LIGHT_AMBER if statistics["pending_reviews"] else LIGHT_BLUE,
    )
    _write_kpi_card(
        worksheet,
        10,
        12,
        "PLAZOS VENCIDOS",
        statistics["overdue_recommendations"],
        "Sin respuesta vigente",
        fill_color=LIGHT_RED if statistics["overdue_recommendations"] else LIGHT_GREEN,
    )
    _write_kpi_card(
        worksheet,
        1,
        15,
        "HALLAZGOS CRÍTICOS",
        statistics["critical_findings"],
        f"{statistics['high_risk_findings']} de riesgo alto o crítico",
        fill_color=LIGHT_RED if statistics["critical_findings"] else LIGHT_BLUE,
    )
    _write_kpi_card(
        worksheet,
        4,
        15,
        "RESOLUCIÓN",
        analysis["resolution_ratio"] if analysis["resolution_ratio"] is not None else "N/D",
        f"De {statistics['total_recommendations']} recomendaciones",
        number_format="0.0%" if analysis["resolution_ratio"] is not None else None,
    )
    worksheet["D15"].comment = Comment(
        "Recomendaciones con resultado definitivo divididas entre el total de recomendaciones.",
        "SIGA-MINEDUCYT",
    )
    _write_kpi_card(
        worksheet,
        7,
        15,
        "INSTITUCIONES",
        statistics["organizations_count"],
        "Incluidas en la selección",
    )
    _write_kpi_card(
        worksheet,
        10,
        15,
        "MEDIANA DE REVISIÓN",
        round(analysis["median_review_days"], 1) if analysis["median_review_days"] is not None else "N/D",
        (
            f"Días corridos · P90 {analysis['p90_review_days']:.1f}"
            if analysis["p90_review_days"] is not None
            else "Sin revisiones resueltas"
        ),
    )

    _write_section_title(worksheet, 19, "Lecturas clave", 12)
    insights = _insights(statistics, analysis, auditor_statistics, organization_statistics)
    for row, insight in enumerate(insights, start=20):
        worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
        cell = worksheet.cell(row=row, column=1, value=f"• {insight}")
        cell.font = Font(name="Aptos", size=10, color="17212B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor=WHITE if row % 2 == 0 else LIGHT_GRAY)
        worksheet.row_dimensions[row].height = 26

    chart_start_row = 22 + len(insights)
    second_chart_row = chart_start_row + 18
    status_rows = statistics["case_status_distribution"]
    recommendation_rows = statistics["recommendation_status_distribution"]
    risk_rows = statistics["risk_distribution"]
    monthly_rows = analysis["series_rows"][-12:]

    data_row = second_chart_row + 20
    worksheet.cell(data_row, 1, "Estado del expediente")
    worksheet.cell(data_row, 2, "Cantidad")
    worksheet.cell(data_row, 3, "Porcentaje")
    for offset, item in enumerate(status_rows, start=1):
        worksheet.cell(data_row + offset, 1, item["label"])
        worksheet.cell(data_row + offset, 2, item["count"])
        worksheet.cell(
            data_row + offset,
            3,
            item["count"] / statistics["total_cases"]
            if statistics["total_cases"]
            else None,
        )

    worksheet.cell(data_row, 5, "Estado de la recomendación")
    worksheet.cell(data_row, 6, "Cantidad")
    worksheet.cell(data_row, 7, "Porcentaje")
    for offset, item in enumerate(recommendation_rows, start=1):
        worksheet.cell(data_row + offset, 5, item["label"])
        worksheet.cell(data_row + offset, 6, item["count"])
        worksheet.cell(
            data_row + offset,
            7,
            item["count"] / statistics["total_recommendations"]
            if statistics["total_recommendations"]
            else None,
        )

    worksheet.cell(data_row, 9, "Nivel de riesgo")
    worksheet.cell(data_row, 10, "Cantidad")
    worksheet.cell(data_row, 11, "Porcentaje")
    for offset, item in enumerate(risk_rows, start=1):
        worksheet.cell(data_row + offset, 9, item["label"])
        worksheet.cell(data_row + offset, 10, item["count"])
        worksheet.cell(
            data_row + offset,
            11,
            item["count"] / statistics["total_findings"]
            if statistics["total_findings"]
            else None,
        )

    monthly_start = max(
        data_row + len(status_rows),
        data_row + len(recommendation_rows),
        data_row + len(risk_rows),
    ) + 3
    worksheet.cell(monthly_start, 1, "Mes de registro")
    worksheet.cell(monthly_start, 2, "Expedientes")
    for offset, item in enumerate(monthly_rows, start=1):
        worksheet.cell(monthly_start + offset, 1, item[1])
        worksheet.cell(monthly_start + offset, 2, item[2])

    for cell_range in (
        f"A{data_row}:C{data_row}",
        f"E{data_row}:G{data_row}",
        f"I{data_row}:K{data_row}",
        f"A{monthly_start}:B{monthly_start}",
    ):
        _style_table_header(worksheet[cell_range][0])
    for column in (3, 7, 11):
        for row in range(data_row + 1, data_row + 1 + max(len(status_rows), len(recommendation_rows), len(risk_rows))):
            worksheet.cell(row, column).number_format = "0.0%"

    if status_rows:
        chart = DoughnutChart()
        chart.title = "Expedientes por estado"
        chart.add_data(
            Reference(worksheet, min_col=2, min_row=data_row, max_row=data_row + len(status_rows)),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(worksheet, min_col=1, min_row=data_row + 1, max_row=data_row + len(status_rows))
        )
        chart.holeSize = 55
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        _add_summary_chart(worksheet, chart, f"A{chart_start_row}")

    if recommendation_rows:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Recomendaciones por resultado"
        chart.add_data(
            Reference(
                worksheet,
                min_col=6,
                min_row=data_row,
                max_row=data_row + len(recommendation_rows),
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                worksheet,
                min_col=5,
                min_row=data_row + 1,
                max_row=data_row + len(recommendation_rows),
            )
        )
        chart.legend = None
        chart.x_axis.title = "Cantidad"
        _add_summary_chart(worksheet, chart, f"G{chart_start_row}")

    if risk_rows:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Hallazgos por nivel de riesgo"
        chart.add_data(
            Reference(worksheet, min_col=10, min_row=data_row, max_row=data_row + len(risk_rows)),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(worksheet, min_col=9, min_row=data_row + 1, max_row=data_row + len(risk_rows))
        )
        chart.legend = None
        chart.y_axis.title = "Cantidad"
        _add_summary_chart(worksheet, chart, f"A{second_chart_row}")

    if monthly_rows:
        chart = LineChart()
        chart.title = "Expedientes registrados por mes"
        chart.add_data(
            Reference(
                worksheet,
                min_col=2,
                min_row=monthly_start,
                max_row=monthly_start + len(monthly_rows),
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                worksheet,
                min_col=1,
                min_row=monthly_start + 1,
                max_row=monthly_start + len(monthly_rows),
            )
        )
        chart.legend = None
        chart.y_axis.title = "Expedientes"
        _add_summary_chart(worksheet, chart, f"G{second_chart_row}")

    focus_start = monthly_start + max(len(monthly_rows), 1) + 3
    _write_section_title(worksheet, focus_start, "Principales focos institucionales", 12)
    focus_headers = (
        "Código",
        "Institución",
        "Expedientes no cerrados",
        "Recomendaciones",
        "Vencidas",
        "Hallazgos críticos",
        "Cumplimiento",
    )
    for column, header in enumerate(focus_headers, start=1):
        worksheet.cell(focus_start + 1, column, header)
    _style_table_header(worksheet[focus_start + 1][: len(focus_headers)])
    for offset, row in enumerate(organization_statistics[:10], start=2):
        values = (
            row["code"],
            row["name"],
            row["open_cases_count"],
            row["recommendation_count"],
            row["overdue_count"],
            row["critical_count"],
            (
                row.get("complied_count", 0) / row.get("terminal_count", 0)
                if row.get("terminal_count", 0)
                else None
            ),
        )
        for column, value in enumerate(values, start=1):
            worksheet.cell(focus_start + offset, column, value)
        worksheet.cell(focus_start + offset, 7).number_format = "0%"

    auditor_start = focus_start + max(len(organization_statistics[:10]), 1) + 4
    _write_section_title(worksheet, auditor_start, "Carga y resultados del equipo auditor", 12)
    auditor_headers = (
        "Auditor",
        "Expedientes no cerrados",
        "Recomendaciones",
        "Por revisar",
        "Vencidas",
        "Cumplimiento",
    )
    for column, header in enumerate(auditor_headers, start=1):
        worksheet.cell(auditor_start + 1, column, header)
    _style_table_header(worksheet[auditor_start + 1][: len(auditor_headers)])
    for offset, row in enumerate(auditor_statistics, start=2):
        values = (
            _display_user(row["auditor"]),
            row["open_cases_count"],
            row["recommendation_count"],
            row["pending_reviews_count"],
            row["overdue_count"],
            (
                row.get("complied_count", 0) / row.get("terminal_count", 0)
                if row.get("terminal_count", 0)
                else None
            ),
        )
        for column, value in enumerate(values, start=1):
            worksheet.cell(auditor_start + offset, column, value)
        worksheet.cell(auditor_start + offset, 6).number_format = "0%"

    worksheet.freeze_panes = "A5"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LETTER
    worksheet.print_title_rows = "1:3"
    worksheet.print_area = f"A1:L{auditor_start + max(len(auditor_statistics), 1) + 2}"
    worksheet.oddFooter.left.text = report_id
    worksheet.oddFooter.center.text = "Página &P de &N"
    worksheet.oddFooter.right.text = "Dirección de Auditoría Interna"
    return worksheet


def _create_metadata_sheet(
    workbook,
    *,
    report_id,
    filters,
    statistics,
    analysis,
    generated_by,
    generated_at,
    counts,
):
    worksheet = workbook.create_sheet(title="Metodología")
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 95
    worksheet.merge_cells("A1:B1")
    worksheet["A1"] = "Metodología, alcance y trazabilidad"
    worksheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    worksheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    worksheet["A1"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 34

    sections = [
        (
            "Identificación",
            [
                ("Folio", report_id),
                ("Fecha y hora de generación", _excel_datetime(generated_at)),
                ("Generado por", _display_user(generated_by)),
                ("Usuario", generated_by.username),
                ("Versión del formato", "1.0"),
            ],
        ),
        (
            "Alcance",
            [
                (
                    "Definición temporal",
                    "El período selecciona expedientes por su fecha de registro. Los estados, "
                    "vencimientos, revisiones y resultados se calculan con la información vigente "
                    "al momento de generar el archivo; no representan una reconstrucción histórica "
                    "al cierre del período.",
                ),
                (
                    "Atribución por auditor",
                    "El auditor corresponde al responsable asignado actualmente al expediente. "
                    "Una reasignación no reconstruye la atribución histórica de actuaciones previas.",
                ),
                (
                    "Instituciones",
                    "La hoja Instituciones agrupa por institución auditada. La hoja Dependencias "
                    "responsables agrupa por la organización responsable de atender cada recomendación.",
                ),
                (
                    "Período",
                    filters.get("selected_period_label", filters["selected_period"]),
                ),
                ("Fecha inicial", filters.get("start_date")),
                ("Fecha final", filters.get("end_date")),
                (
                    "Auditor",
                    _display_user(filters.get("selected_auditor")) or "Todos",
                ),
                (
                    "Institución",
                    filters["selected_organization"].name
                    if filters.get("selected_organization")
                    else "Todas",
                ),
                (
                    "Estado",
                    dict(AuditCase.Status.choices).get(
                        filters.get("selected_status"), "Todos"
                    ),
                ),
            ],
        ),
        (
            "Conteos exportados",
            [
                ("Expedientes", counts["cases"]),
                ("Hallazgos", counts["findings"]),
                ("Recomendaciones", counts["recommendations"]),
                ("Respuestas", counts["responses"]),
                ("Prórrogas", counts["extensions"]),
                ("Alertas y controles", counts["alerts"]),
                ("Auditores", counts["auditors"]),
                ("Instituciones auditadas", counts["audited_organizations"]),
                ("Dependencias responsables", counts["responsible_organizations"]),
                ("Meses en la serie", counts["series_months"]),
                ("Reconciliaciones", counts["controls"]),
            ],
        ),
        (
            "Definiciones de indicadores",
            [
                (
                    "Cumplimiento",
                    "Recomendaciones calificadas como cumplidas divididas entre las que poseen "
                    "un resultado definitivo: cumplida, parcialmente cumplida o no cumplida.",
                ),
                (
                    "Resolución",
                    "Recomendaciones con resultado definitivo divididas entre el total de recomendaciones.",
                ),
                (
                    "Plazo vencido",
                    "Recomendación pendiente o que requiere corrección cuya fecha límite vigente, "
                    "incluida la última prórroga concedida, es anterior a la fecha de generación.",
                ),
                (
                    "Incumplimiento automático",
                    "Recomendación marcada como no cumplida por falta de respuesta al ejecutarse "
                    "el proceso automático de vencimientos. Ya no se incluye entre los plazos "
                    "vencidos aún pendientes.",
                ),
                (
                    "Versión de respuesta por revisar",
                    "Cada versión de respuesta presentada que aún no posee un dictamen de Auditoría. "
                    "Una recomendación puede tener más de una versión a lo largo del seguimiento.",
                ),
                (
                    "Tiempo de revisión",
                    "Días corridos transcurridos entre la presentación de cada versión de respuesta "
                    "y su revisión. El resumen muestra mediana y percentil 90; el detalle conserva "
                    "también las horas corridas.",
                ),
                (
                    "No cerrados y en gestión",
                    "No cerrados incluye todos los estados excepto Cerrado, incluidos los borradores. "
                    "En gestión excluye tanto Borrador como Cerrado.",
                ),
                (
                    "Focos institucionales",
                    "Ordenados primero por recomendaciones vencidas, luego por hallazgos críticos "
                    "y finalmente por expedientes activos. No constituye una calificación de riesgo formal.",
                ),
            ],
        ),
        (
            "Indicadores complementarios del corte",
            [
                ("Expedientes en gestión", analysis["cases_in_management"]),
                ("Cumplimiento parcial", analysis["partial_count"]),
                ("No cumplidas", analysis["not_complied_count"]),
                (
                    "Incumplimientos automáticos por falta de respuesta",
                    analysis["automatic_non_compliance_count"],
                ),
                ("Próximas a vencer en 7 días", analysis["due_soon_count"]),
                ("Recomendaciones sin fecha límite", analysis["no_deadline_count"]),
                ("Versiones de respuesta pendientes de revisión", analysis["pending_response_versions"]),
                ("Recomendaciones con versiones pendientes", analysis["pending_recommendations"]),
                ("Mediana de revisión en días corridos", analysis["median_review_days"]),
                ("Percentil 90 de revisión en días corridos", analysis["p90_review_days"]),
            ],
        ),
        (
            "Seguridad y uso",
            [
                (
                    "Clasificación",
                    "Información interna de la Dirección de Auditoría. Debe conservarse y compartirse "
                    "según las políticas institucionales aplicables.",
                ),
                (
                    "Datos excluidos",
                    "El libro no incluye archivos de evidencia, rutas privadas de almacenamiento, "
                    "credenciales ni huellas de los documentos cargados.",
                ),
                (
                    "Integridad",
                    "La descarga queda registrada en la bitácora del sistema junto con los filtros, "
                    "conteos y huella SHA-256 del archivo.",
                ),
            ],
        ),
    ]

    row = 3
    for section, values in sections:
        worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = worksheet.cell(row, 1, section)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center")
        row += 1
        for label, value in values:
            worksheet.cell(row, 1, label)
            value_cell = worksheet.cell(row, 2, value)
            worksheet.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_GRAY)
            worksheet.cell(row, 1).font = Font(name="Aptos", size=9, bold=True, color=NAVY)
            value_cell.font = Font(name="Aptos", size=9, color="17212B")
            value_cell.alignment = Alignment(vertical="top", wrap_text=True)
            worksheet.cell(row, 1).alignment = Alignment(vertical="top", wrap_text=True)
            for column in (1, 2):
                worksheet.cell(row, column).border = Border(bottom=Side(style="hair", color=LINE))
            if isinstance(value, date) and not isinstance(value, datetime):
                value_cell.number_format = "dd/mm/yyyy"
            elif isinstance(value, datetime):
                value_cell.number_format = "dd/mm/yyyy hh:mm"
            worksheet.row_dimensions[row].height = 32 if len(str(value or "")) > 100 else 22
            row += 1
        row += 1

    worksheet.freeze_panes = "A3"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_setup.orientation = "portrait"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LETTER
    worksheet.print_area = f"A1:B{row}"
    return worksheet


def build_director_statistics_xlsx(
    *,
    cases,
    filters,
    statistics,
    auditor_statistics,
    organization_statistics,
    generated_by,
    generated_at=None,
):
    generated_at = generated_at or timezone.now()
    today = timezone.localdate(generated_at)
    report_id = f"DAI-XLSX-{timezone.localtime(generated_at):%Y%m%d-%H%M%S}-{generated_by.pk}"
    dataset = _build_dataset(cases, today)
    analysis = _extended_analysis(dataset, filters, today)
    alert_rows = _alert_rows(dataset, today)
    responsible_organization_rows = _responsible_organization_rows(dataset, today)
    control_rows = _control_rows(
        dataset, statistics, organization_statistics, alert_rows
    )

    workbook = Workbook()
    workbook.properties.creator = _display_user(generated_by)
    workbook.properties.lastModifiedBy = _display_user(generated_by)
    workbook.properties.title = "Informe estadístico y analítico de Auditoría Interna"
    workbook.properties.subject = report_id
    workbook.properties.description = (
        "Corte estadístico generado por SIGA-MINEDUCYT."
    )
    workbook.properties.keywords = "auditoría, estadísticas, recomendaciones, cumplimiento, riesgo"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    _create_summary_sheet(
        workbook,
        report_id=report_id,
        filters=filters,
        statistics=statistics,
        analysis=analysis,
        auditor_statistics=auditor_statistics,
        organization_statistics=organization_statistics,
        generated_by=generated_by,
        generated_at=generated_at,
    )

    case_rows = _case_rows(dataset)
    case_sheet = _create_data_sheet(
        workbook,
        title="Expedientes",
        description=(
            "Expedientes incluidos en el corte, con indicadores operativos calculados sobre sus "
            "hallazgos, recomendaciones y respuestas."
        ),
        headers=(
            "Referencia",
            "Título",
            "Código institución",
            "Institución auditada",
            "Departamento",
            "Municipio",
            "Auditor responsable actual",
            "Usuario auditor",
            "Estado",
            "Fecha del informe",
            "Inicio período auditado",
            "Fin período auditado",
            "Fecha límite general",
            "Fecha de registro",
            "Última actualización",
            "Hallazgos",
            "Hallazgos altos/críticos",
            "Recomendaciones",
            "Recomendaciones vencidas",
            "Respuestas",
            "Versiones de respuesta por revisar",
            "Cumplimiento",
        ),
        rows=case_rows,
        table_name="TablaExpedientes",
        widths={1: 18, 2: 38, 3: 18, 4: 38, 5: 18, 6: 20, 7: 25, 8: 18, 9: 22},
        date_columns=(10, 11, 12, 13),
        datetime_columns=(14, 15),
        integer_columns=(16, 17, 18, 19, 20, 21),
        percentage_columns=(22,),
        wrap_columns=(2, 4),
    )
    if case_rows:
        case_sheet.conditional_formatting.add(
            f"S5:S{4 + len(case_rows)}",
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill("solid", fgColor=LIGHT_RED),
                font=Font(color=RED, bold=True),
            ),
        )

    finding_rows = _finding_rows(dataset)
    finding_sheet = _create_data_sheet(
        workbook,
        title="Hallazgos",
        description="Detalle completo de los hallazgos y sus elementos de análisis.",
        headers=(
            "Expediente",
            "Código institución",
            "Institución auditada",
            "Auditor",
            "N.º hallazgo",
            "Título",
            "Nivel de riesgo",
            "Condición",
            "Criterio",
            "Causa",
            "Efecto",
            "Recomendaciones",
        ),
        rows=finding_rows,
        table_name="TablaHallazgos",
        widths={1: 18, 2: 18, 3: 35, 4: 24, 5: 12, 6: 38, 7: 16, 8: 48, 9: 48, 10: 48, 11: 48},
        integer_columns=(5, 12),
        wrap_columns=(3, 6, 8, 9, 10, 11),
    )
    if finding_rows:
        risk_range = f"G5:G{4 + len(finding_rows)}"
        finding_sheet.conditional_formatting.add(
            risk_range,
            FormulaRule(
                formula=['$G5="Crítico"'],
                fill=PatternFill("solid", fgColor=LIGHT_RED),
                font=Font(color=RED, bold=True),
            ),
        )
        finding_sheet.conditional_formatting.add(
            risk_range,
            FormulaRule(
                formula=['$G5="Alto"'],
                fill=PatternFill("solid", fgColor=LIGHT_AMBER),
                font=Font(color=AMBER, bold=True),
            ),
        )

    recommendation_rows = _recommendation_rows(dataset, today)
    recommendation_sheet = _create_data_sheet(
        workbook,
        title="Recomendaciones",
        description=(
            "Seguimiento detallado de cada recomendación. La fecha límite vigente incorpora la "
            "última prórroga concedida."
        ),
        headers=(
            "Expediente",
            "Código institución",
            "Institución auditada",
            "Auditor",
            "N.º hallazgo",
            "Hallazgo",
            "Riesgo",
            "N.º recomendación",
            "Recomendación",
            "Código responsable",
            "Institución responsable",
            "Estado",
            "Fecha límite original",
            "Fecha límite vigente",
            "Días vencidos",
            "Días para vencer",
            "Respuestas",
            "Prórrogas",
            "Última respuesta",
            "Estado declarado",
            "Último dictamen",
            "Fecha del dictamen",
            "Revisada por",
            "Evidencias requeridas",
            "Procedencia",
            "Incumplimiento automático registrado",
        ),
        rows=recommendation_rows,
        table_name="TablaRecomendaciones",
        widths={1: 18, 2: 18, 3: 35, 4: 24, 6: 36, 9: 55, 10: 18, 11: 35, 12: 23, 19: 22, 20: 23, 22: 24, 23: 48, 24: 35},
        date_columns=(13, 14),
        datetime_columns=(19, 22, 26),
        integer_columns=(5, 8, 15, 16, 17, 18),
        wrap_columns=(3, 6, 9, 11, 24, 25),
    )
    if recommendation_rows:
        last_row = 4 + len(recommendation_rows)
        recommendation_sheet.conditional_formatting.add(
            f"O5:O{last_row}",
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill("solid", fgColor=LIGHT_RED),
                font=Font(color=RED, bold=True),
            ),
        )
        recommendation_sheet.conditional_formatting.add(
            f"P5:P{last_row}",
            CellIsRule(
                operator="between",
                formula=["0", "7"],
                fill=PatternFill("solid", fgColor=LIGHT_AMBER),
                font=Font(color=AMBER, bold=True),
            ),
        )
        _add_status_conditional_formatting(recommendation_sheet, "L", 5, last_row)

    response_rows = _response_rows(dataset)
    response_sheet = _create_data_sheet(
        workbook,
        title="Respuestas y revisiones",
        description=(
            "Versiones de respuestas institucionales, acciones declaradas, cantidad de evidencias "
            "y dictamen de Auditoría cuando existe."
        ),
        headers=(
            "Expediente",
            "Código institución",
            "Institución auditada",
            "Auditor",
            "N.º hallazgo",
            "N.º recomendación",
            "Versión",
            "Estado declarado",
            "Acciones realizadas",
            "Fecha de ejecución",
            "Responsable",
            "Cargo",
            "Motivo de incumplimiento",
            "Plan de acción",
            "Finalización estimada",
            "Declaración de veracidad",
            "Presentada por",
            "Fecha de presentación",
            "Evidencias",
            "Evidencias pendientes de análisis",
            "Evidencias aprobadas",
            "Evidencias rechazadas",
            "CDE registrado",
            "Dictamen de Auditoría",
            "Comentarios de Auditoría",
            "Revisada por",
            "Fecha de revisión",
            "Horas corridas de revisión",
            "Días corridos de revisión",
        ),
        rows=response_rows,
        table_name="TablaRespuestas",
        widths={1: 18, 2: 18, 3: 35, 4: 24, 8: 22, 9: 55, 11: 28, 12: 28, 13: 45, 14: 45, 16: 25, 19: 18, 20: 24, 21: 48, 22: 25},
        date_columns=(10, 15),
        datetime_columns=(18, 27),
        integer_columns=(5, 6, 7, 19, 20, 21, 22),
        wrap_columns=(3, 9, 11, 12, 13, 14, 25),
    )
    if response_rows:
        _add_status_conditional_formatting(response_sheet, "X", 5, 4 + len(response_rows))

    extension_rows = _extension_rows(dataset)
    _create_data_sheet(
        workbook,
        title="Prórrogas",
        description=(
            "Historial de prórrogas concedidas. Cada fila conserva el plazo anterior, los días "
            "hábiles otorgados y el nuevo vencimiento."
        ),
        headers=(
            "Expediente",
            "Código institución",
            "Institución auditada",
            "Auditor",
            "N.º hallazgo",
            "N.º recomendación",
            "Fecha límite anterior",
            "Días hábiles",
            "Nueva fecha límite",
            "Motivo",
            "Concedida por",
            "Fecha de registro",
        ),
        rows=extension_rows,
        table_name="TablaProrrogas",
        widths={1: 18, 2: 18, 3: 35, 4: 24, 7: 20, 8: 16, 9: 20, 10: 55, 11: 25},
        date_columns=(7, 9),
        datetime_columns=(12,),
        integer_columns=(5, 6, 8),
        wrap_columns=(3, 10),
    )

    auditor_rows = _auditor_rows(auditor_statistics)
    auditor_sheet = _create_data_sheet(
        workbook,
        title="Auditores",
        description=(
            "Comparativo atribuido al auditor responsable actual de cada expediente. Una reasignación "
            "no reconstruye la responsabilidad histórica."
        ),
        headers=(
            "Auditor",
            "Usuario",
            "Cargo",
            "Activo",
            "Expedientes en la selección",
            "Expedientes no cerrados",
            "Expedientes en gestión",
            "Recomendaciones",
            "Cumplidas",
            "Parcialmente cumplidas",
            "No cumplidas",
            "Versiones de respuesta por revisar",
            "Recomendaciones vencidas",
            "Cumplimiento",
        ),
        rows=auditor_rows,
        table_name="TablaAuditores",
        widths={1: 28, 2: 20, 3: 28, 4: 12, 5: 24, 6: 22, 7: 22, 8: 18, 9: 16, 10: 22, 11: 18, 12: 22, 13: 24, 14: 18},
        integer_columns=(5, 6, 7, 8, 9, 10, 11, 12, 13),
        percentage_columns=(14,),
    )
    if auditor_rows:
        auditor_sheet.conditional_formatting.add(
            f"M5:M{4 + len(auditor_rows)}",
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill("solid", fgColor=LIGHT_RED),
                font=Font(color=RED, bold=True),
            ),
        )

    organization_rows = _organization_rows(organization_statistics)
    organization_sheet = _create_data_sheet(
        workbook,
        title="Instituciones",
        description=(
            "Análisis completo por institución, ordenado por vencimientos, hallazgos críticos y "
            "expedientes no cerrados."
        ),
        headers=(
            "Código",
            "Institución",
            "Tipo",
            "Departamento",
            "Municipio",
            "Expedientes en la selección",
            "Expedientes no cerrados",
            "Expedientes en gestión",
            "Recomendaciones",
            "Cumplidas",
            "Parcialmente cumplidas",
            "No cumplidas",
            "Recomendaciones vencidas",
            "Hallazgos críticos",
            "Cumplimiento",
        ),
        rows=organization_rows,
        table_name="TablaInstituciones",
        widths={1: 18, 2: 42, 3: 24, 4: 20, 5: 22, 6: 24, 7: 22, 8: 22, 9: 18, 10: 16, 11: 22, 12: 18, 13: 24, 14: 20, 15: 18},
        integer_columns=(6, 7, 8, 9, 10, 11, 12, 13, 14),
        percentage_columns=(15,),
        wrap_columns=(2,),
    )
    if organization_rows:
        organization_sheet.conditional_formatting.add(
            f"M5:M{4 + len(organization_rows)}",
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill("solid", fgColor=LIGHT_RED),
                font=Font(color=RED, bold=True),
            ),
        )

    responsible_sheet = _create_data_sheet(
        workbook,
        title="Dependencias responsables",
        description=(
            "Agrupación por la organización responsable de atender cada recomendación. Puede ser "
            "distinta de la institución auditada."
        ),
        headers=(
            "Código",
            "Dependencia responsable",
            "Tipo",
            "Departamento",
            "Municipio",
            "Expedientes relacionados",
            "Recomendaciones",
            "Pendientes",
            "Respuesta enviada",
            "En revisión",
            "Requieren corrección",
            "Cumplidas",
            "Parcialmente cumplidas",
            "No cumplidas",
            "Cumplimiento",
            "Vencidas pendientes",
            "Próximas a vencer",
            "Sin fecha límite",
            "Incumplimientos automáticos",
        ),
        rows=responsible_organization_rows,
        table_name="TablaDependenciasResponsables",
        widths={1: 18, 2: 42, 3: 24, 4: 20, 5: 22, 6: 22, 7: 18, 8: 16, 9: 20, 10: 16, 11: 22, 12: 16, 13: 22, 14: 18, 15: 18, 16: 20, 17: 20, 18: 18, 19: 26},
        integer_columns=tuple(range(6, 15)) + tuple(range(16, 20)),
        percentage_columns=(15,),
        wrap_columns=(2,),
    )
    if responsible_organization_rows:
        responsible_sheet.conditional_formatting.add(
            f"P5:P{4 + len(responsible_organization_rows)}",
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill("solid", fgColor=LIGHT_RED),
                font=Font(color=RED, bold=True),
            ),
        )

    series_sheet = _create_data_sheet(
        workbook,
        title="Series mensuales",
        description=(
            "Serie continua del conjunto de expedientes seleccionado. Las respuestas y revisiones "
            "pueden haber ocurrido fuera del período de registro del expediente."
        ),
        headers=(
            "Mes",
            "Etiqueta",
            "Expedientes registrados",
            "Respuestas presentadas",
            "Revisiones emitidas",
        ),
        rows=analysis["series_rows"],
        table_name="TablaSeriesMensuales",
        widths={1: 16, 2: 16, 3: 24, 4: 24, 5: 22},
        date_columns=(1,),
        integer_columns=(3, 4, 5),
        landscape=False,
    )
    series_sheet.freeze_panes = "A5"

    alert_sheet = _create_data_sheet(
        workbook,
        title="Alertas y calidad",
        description=(
            "Excepciones operativas y controles de completitud. Las alertas apoyan la priorización "
            "y no sustituyen el criterio profesional de Auditoría."
        ),
        headers=(
            "Categoría",
            "Prioridad",
            "Alerta o control",
            "Identificador",
            "Institución",
            "Auditor",
            "Fecha límite",
            "Días",
            "Detalle",
        ),
        rows=alert_rows,
        table_name="TablaAlertas",
        widths={1: 22, 2: 14, 3: 35, 4: 30, 5: 38, 6: 25, 7: 18, 8: 12, 9: 60},
        date_columns=(7,),
        integer_columns=(8,),
        wrap_columns=(3, 5, 9),
    )
    if alert_rows:
        last_row = 4 + len(alert_rows)
        for label, color, font_color in (
            ("Crítica", LIGHT_RED, RED),
            ("Alta", LIGHT_AMBER, AMBER),
            ("Media", LIGHT_BLUE, BLUE),
        ):
            alert_sheet.conditional_formatting.add(
                f"B5:B{last_row}",
                FormulaRule(
                    formula=[f'$B5="{label}"'],
                    fill=PatternFill("solid", fgColor=color),
                    font=Font(color=font_color, bold=True),
                ),
            )

    control_sheet = _create_data_sheet(
        workbook,
        title="Controles",
        description=(
            "Reconciliaciones internas del libro. Todos los controles deben indicar OK al momento "
            "de la generación."
        ),
        headers=("Control", "Valor esperado", "Valor exportado", "Resultado", "Descripción"),
        rows=control_rows,
        table_name="TablaControles",
        widths={1: 38, 2: 18, 3: 18, 4: 14, 5: 65},
        integer_columns=(2, 3),
        wrap_columns=(1, 5),
        landscape=False,
    )
    if control_rows:
        control_sheet.conditional_formatting.add(
            f"D5:D{4 + len(control_rows)}",
            FormulaRule(
                formula=['$D5="ERROR"'],
                fill=PatternFill("solid", fgColor=LIGHT_RED),
                font=Font(color=RED, bold=True),
            ),
        )
        control_sheet.conditional_formatting.add(
            f"D5:D{4 + len(control_rows)}",
            FormulaRule(
                formula=['$D5="OK"'],
                fill=PatternFill("solid", fgColor=LIGHT_GREEN),
                font=Font(color=GREEN, bold=True),
            ),
        )

    counts = {
        "cases": len(dataset["cases"]),
        "findings": len(dataset["findings"]),
        "recommendations": len(dataset["recommendations"]),
        "responses": len(dataset["responses"]),
        "extensions": len(dataset["extensions"]),
        "alerts": len(alert_rows),
        "auditors": len(auditor_rows),
        "audited_organizations": len(organization_rows),
        "responsible_organizations": len(responsible_organization_rows),
        "series_months": len(analysis["series_rows"]),
        "controls": len(control_rows),
    }
    _create_metadata_sheet(
        workbook,
        report_id=report_id,
        filters=filters,
        statistics=statistics,
        analysis=analysis,
        generated_by=generated_by,
        generated_at=generated_at,
        counts=counts,
    )

    workbook.active = 0
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer, {"report_id": report_id, "counts": counts}
