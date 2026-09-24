from datetime import date, timedelta
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.institutions.models import Organization

from .models import (
    ActivityLog,
    AuditCase,
    DeadlineExtension,
    Finding,
    Recommendation,
    Response,
    Review,
)


class TerritorialAnalysisTests(TestCase):
    def setUp(self):
        self.audit_unit = Organization.objects.create(
            code="DAI-TERR",
            name="Dirección de Auditoría",
            kind=Organization.Kind.MINISTRY_UNIT,
        )
        self.center_a = Organization.objects.create(
            code="CE-A",
            name="Centro Escolar A",
            kind=Organization.Kind.EDUCATIONAL_CENTER,
            department="Departamento A",
            municipality="Municipio A",
            district="Distrito A1",
        )
        self.center_b = Organization.objects.create(
            code="CE-B",
            name="Centro Escolar B",
            kind=Organization.Kind.EDUCATIONAL_CENTER,
            department="Departamento B",
            municipality="Municipio B",
            district="Distrito B1",
        )
        self.center_without_audits = Organization.objects.create(
            code="CE-C",
            name="Centro Escolar sin auditorías",
            kind=Organization.Kind.EDUCATIONAL_CENTER,
            department="Departamento A",
            municipality="Municipio A",
            district="Distrito A2",
        )
        self.director = User.objects.create_user(
            username="directora-territorial",
            password="ClaveSegura!2026",
            role=User.Role.AUDIT_MANAGER,
            organization=self.audit_unit,
            must_change_password=False,
        )
        self.auditor = User.objects.create_user(
            username="auditor-territorial",
            password="ClaveSegura!2026",
            role=User.Role.AUDITOR,
            organization=self.audit_unit,
            must_change_password=False,
        )

    def _recommendation(
        self,
        *,
        audited,
        responsible,
        reference,
        risk=Finding.RiskLevel.MEDIUM,
        status=Recommendation.Status.PENDING,
        deadline=None,
        case_status=AuditCase.Status.PUBLISHED,
    ):
        case = AuditCase.objects.create(
            reference=reference,
            title=f"Expediente {reference}",
            audited_organization=audited,
            assigned_auditor=self.auditor,
            created_by=self.auditor,
            status=case_status,
            report_date=date.today(),
        )
        finding = Finding.objects.create(
            case=case,
            number=1,
            title=f"Hallazgo {reference}",
            risk_level=risk,
        )
        recommendation = Recommendation.objects.create(
            finding=finding,
            number=1,
            text="Atender la condición observada.",
            responsible_organization=responsible,
            status=status,
            deadline=deadline,
        )
        return case, recommendation

    def test_analysis_is_exclusive_to_director(self):
        url = reverse("director_statistics")

        anonymous = self.client.get(url)
        self.assertRedirects(anonymous, f"{reverse('login')}?next={url}")

        self.client.force_login(self.auditor)
        self.assertEqual(self.client.get(url).status_code, 403)

        self.client.force_login(self.director)
        allowed = self.client.get(url)
        self.assertEqual(allowed.status_code, 200)
        self.assertContains(allowed, "Análisis de centros educativos")

    def test_secondary_filters_are_collapsed_until_one_is_active(self):
        self.client.force_login(self.director)
        url = reverse("director_statistics")

        default = self.client.get(url)
        filtered = self.client.get(
            url,
            {"mode": "activity", "coverage": "not_audited", "period": "30"},
        )

        self.assertEqual(default.context["advanced_filter_count"], 0)
        self.assertFalse(default.context["advanced_filters_open"])
        self.assertContains(default, '<details class="territorial-advanced-filters" >')
        self.assertEqual(filtered.context["advanced_filter_count"], 1)
        self.assertTrue(filtered.context["advanced_filters_open"])
        self.assertContains(
            filtered, '<details class="territorial-advanced-filters" open>'
        )
        self.assertContains(filtered, "1 activo")

    def test_custom_period_errors_are_exposed_in_the_period_selector(self):
        self.client.force_login(self.director)

        response = self.client.get(
            reverse("director_statistics"),
            {"mode": "activity", "period": "custom"},
        )

        self.assertTrue(response.context["filter_errors"])
        self.assertFalse(response.context["advanced_filters_open"])
        self.assertContains(
            response,
            '<fieldset class="period-filter-group analysis-period-selector" data-period-filter >',
        )
        self.assertContains(response, "Indique una fecha inicial y una fecha final válidas.")

    def test_current_mode_does_not_validate_the_hidden_activity_period(self):
        self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-CURRENT-CUSTOM",
        )
        self.client.force_login(self.director)

        response = self.client.get(
            reverse("director_statistics"),
            {"mode": "current", "period": "custom"},
        )

        self.assertEqual(response.context["filter_errors"], [])
        self.assertEqual(response.context["audited_center_count"], 1)
        self.assertContains(response, "Desde todos los tiempos")
        self.assertNotContains(
            response, "Indique una fecha inicial y una fecha final válidas."
        )

    def test_valid_custom_period_filters_activity_and_is_kept_in_mode_links(self):
        included_case, _recommendation = self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-CUSTOM-IN",
        )
        excluded_case, _recommendation = self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-CUSTOM-OUT",
        )
        AuditCase.objects.filter(pk=excluded_case.pk).update(
            created_at=timezone.now() - timedelta(days=3)
        )
        self.client.force_login(self.director)
        selected_date = date.today().isoformat()

        response = self.client.get(
            reverse("director_statistics"),
            {
                "mode": "activity",
                "period": "custom",
                "start": selected_date,
                "end": selected_date,
                "department": "Departamento A",
            },
        )
        content = response.content.decode()

        self.assertEqual(response.context["filter_errors"], [])
        self.assertEqual(response.context["activity_case_count"], 1)
        self.assertEqual(response.context["activity_finding_count"], 1)
        self.assertEqual(response.context["activity_recommendation_count"], 1)
        self.assertEqual(response.context["selected_period"], "custom")
        self.assertIn(
            f"mode=current&amp;period=custom&amp;start={selected_date}&amp;end={selected_date}",
            content,
        )
        self.assertContains(
            response,
            '<div class="period-custom-dates" data-custom-period-dates >',
        )
        self.assertContains(response, f'value="{included_case.created_at:%Y-%m-%d}"')
        self.assertNotContains(response, 'id="current-title"')

    def test_quick_filters_prioritize_actionable_audit_signals(self):
        self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-QUICK-OVERDUE",
            deadline=date.today() - timedelta(days=2),
        )
        self._recommendation(
            audited=self.center_b,
            responsible=self.center_b,
            reference="IA-TERR-QUICK-CRITICAL",
            risk=Finding.RiskLevel.CRITICAL,
            status=Recommendation.Status.NOT_COMPLIED,
        )
        self.client.force_login(self.director)

        response = self.client.get(reverse("director_statistics"))
        quick_filters = {
            item["value"]: item for item in response.context["quick_filters"]
        }

        self.assertEqual(quick_filters["immediate"]["count"], 2)
        self.assertEqual(quick_filters["overdue"]["count"], 1)
        self.assertEqual(quick_filters["critical"]["count"], 1)
        self.assertEqual(quick_filters["not_complied"]["count"], 1)
        self.assertEqual(quick_filters["never_audited"]["count"], 1)
        self.assertNotIn("cde", quick_filters)
        self.assertContains(response, "Prioridades")
        self.assertContains(response, "Plazos vencidos")
        self.assertContains(response, "Nunca auditados")

    def test_quick_filter_limits_centers_and_preserves_geographic_scope(self):
        self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-QUICK-SCOPE",
            deadline=date.today() - timedelta(days=2),
        )
        self.client.force_login(self.director)

        response = self.client.get(
            reverse("director_statistics"),
            {"department": "Departamento A", "quick": "overdue"},
        )

        self.assertEqual(response.context["selected_quick_filter"], "overdue")
        self.assertEqual(response.context["center_count"], 1)
        self.assertEqual(response.context["centers"], [self.center_a])
        overdue_filter = next(
            item
            for item in response.context["quick_filters"]
            if item["value"] == "overdue"
        )
        self.assertTrue(overdue_filter["is_active"])
        self.assertIn("department=Departamento+A", overdue_filter["url"])
        self.assertContains(response, 'name="quick" value="overdue"')

    def test_universe_includes_never_audited_centers_and_groups_by_territory(self):
        self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-001",
        )
        self.client.force_login(self.director)

        response = self.client.get(
            reverse("director_statistics"), {"group_by": "department"}
        )

        self.assertEqual(response.context["center_count"], 3)
        self.assertEqual(response.context["audited_center_count"], 1)
        self.assertEqual(response.context["not_audited_center_count"], 2)
        department_a = next(
            row
            for row in response.context["territory_rows"]
            if row["name"] == "Departamento A"
        )
        self.assertEqual(department_a["center_count"], 2)
        self.assertEqual(department_a["audited_count"], 1)
        self.assertEqual(department_a["coverage_rate"], 50)

    def test_district_options_and_results_are_limited_to_selected_department(self):
        self.client.force_login(self.director)
        url = reverse("director_statistics")

        district = self.client.get(
            url,
            {
                "department": "Departamento A",
                "district": "Distrito A2",
                "group_by": "district",
            },
        )
        invalid = self.client.get(
            url,
            {
                "department": "Departamento A",
                "district": "Distrito B1",
            },
        )

        self.assertEqual(district.context["district_options"], ["Distrito A1", "Distrito A2"])
        self.assertNotContains(district, 'name="municipality"')
        self.assertNotContains(district, "Distrito B1")
        self.assertEqual(district.context["center_count"], 1)
        self.assertEqual(district.context["centers"][0], self.center_without_audits)
        self.assertEqual(invalid.context["center_count"], 0)
        self.assertNotContains(invalid, self.center_a.name)

    def test_center_directory_uses_correlated_district_filter(self):
        self.client.force_login(self.director)

        response = self.client.get(
            reverse("director_educational_centers"),
            {"department": "Departamento A", "district": "Distrito A2"},
        )

        self.assertEqual(list(response.context["district_options"]), ["Distrito A1", "Distrito A2"])
        self.assertNotContains(response, 'name="municipality"')
        self.assertContains(response, self.center_without_audits.name)
        self.assertNotContains(response, self.center_a.name)
        self.assertNotContains(response, self.center_b.name)

    def test_old_overdue_obligation_remains_in_current_snapshot_with_30_day_period(self):
        case, _recommendation = self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-OLD",
            deadline=date.today() - timedelta(days=5),
        )
        AuditCase.objects.filter(pk=case.pk).update(
            created_at=timezone.now() - timedelta(days=100)
        )
        self.client.force_login(self.director)

        response = self.client.get(
            reverse("director_statistics"),
            {"period": "30", "department": "Departamento A"},
        )

        self.assertEqual(response.context["overdue_center_count"], 1)
        self.assertEqual(response.context["overdue_obligations"], 1)
        self.assertEqual(response.context["activity_case_count"], 0)

    def test_risk_and_overdue_are_attributed_to_different_territories(self):
        self._recommendation(
            audited=self.center_a,
            responsible=self.center_b,
            reference="IA-TERR-SPLIT",
            risk=Finding.RiskLevel.CRITICAL,
            deadline=date.today() - timedelta(days=2),
        )
        self.client.force_login(self.director)
        url = reverse("director_statistics")

        audited_territory = self.client.get(
            url, {"department": "Departamento A"}
        )
        responsible_territory = self.client.get(
            url, {"department": "Departamento B"}
        )

        self.assertEqual(audited_territory.context["critical_center_count"], 1)
        self.assertEqual(audited_territory.context["overdue_center_count"], 0)
        self.assertEqual(responsible_territory.context["critical_center_count"], 0)
        self.assertEqual(responsible_territory.context["overdue_center_count"], 1)

    def test_drafts_do_not_change_current_or_period_indicators(self):
        self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-DRAFT",
            risk=Finding.RiskLevel.CRITICAL,
            deadline=date.today() - timedelta(days=30),
            case_status=AuditCase.Status.DRAFT,
        )
        self.client.force_login(self.director)

        response = self.client.get(
            reverse("director_statistics"), {"department": "Departamento A"}
        )

        self.assertEqual(response.context["audited_center_count"], 0)
        self.assertEqual(response.context["critical_center_count"], 0)
        self.assertEqual(response.context["overdue_center_count"], 0)
        self.assertEqual(response.context["activity_case_count"], 0)

    def test_territorial_compliance_uses_aggregate_denominator_and_not_available(self):
        _, first = self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-COMP-1",
            status=Recommendation.Status.COMPLIED,
        )
        case = first.finding.case
        Recommendation.objects.create(
            finding=case.findings.get(),
            number=2,
            text="Segunda recomendación.",
            responsible_organization=self.center_a,
            status=Recommendation.Status.NOT_COMPLIED,
        )
        self._recommendation(
            audited=self.center_b,
            responsible=self.center_b,
            reference="IA-TERR-COMP-2",
            status=Recommendation.Status.COMPLIED,
        )
        self.client.force_login(self.director)

        country = self.client.get(reverse("director_statistics"))
        no_results = self.client.get(
            reverse("director_statistics"), {"district": "Distrito A2"}
        )

        self.assertEqual(country.context["terminal_obligations"], 3)
        self.assertEqual(country.context["complied_obligations"], 2)
        self.assertEqual(country.context["territorial_compliance_rate"], 67)
        self.assertIsNone(no_results.context["territorial_compliance_rate"])
        self.assertContains(no_results, "N/D")

    def test_activity_period_uses_each_event_date_and_effective_extension(self):
        _case, recommendation = self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-ACTIVITY",
            deadline=date.today() - timedelta(days=1),
        )
        DeadlineExtension.objects.create(
            recommendation=recommendation,
            previous_deadline=recommendation.deadline,
            business_days=5,
            new_deadline=date.today() + timedelta(days=5),
            reason="Prórroga territorial de prueba.",
            granted_by=self.auditor,
        )
        response = Response.objects.create(
            recommendation=recommendation,
            version=1,
            declared_status=Response.DeclaredStatus.IN_PROGRESS,
            action_description="Acciones iniciadas.",
            responsible_name="Responsable",
            responsible_job_title="Dirección",
            accuracy_declaration=True,
            submitted_by=self.auditor,
        )
        Review.objects.create(
            response=response,
            outcome=Review.Outcome.PARTIAL,
            comments="Cumplimiento parcial.",
            reviewed_by=self.auditor,
        )
        self.client.force_login(self.director)

        result = self.client.get(
            reverse("director_statistics"),
            {"mode": "activity", "period": "30", "department": "Departamento A"},
        )

        self.assertEqual(result.context["activity_case_count"], 1)
        self.assertEqual(result.context["activity_report_count"], 1)
        self.assertEqual(result.context["activity_response_count"], 1)
        self.assertEqual(result.context["activity_review_count"], 1)
        self.assertEqual(result.context["activity_extension_count"], 1)
        self.assertEqual(result.context["overdue_center_count"], 0)
        self.assertEqual(result.context["centers"][0].due_soon_count, 1)

    def test_xlsx_uses_the_same_territorial_filters_and_metrics_as_the_page(self):
        self._recommendation(
            audited=self.center_a,
            responsible=self.center_a,
            reference="IA-TERR-XLSX-IN",
            deadline=date.today() - timedelta(days=2),
        )
        self._recommendation(
            audited=self.center_b,
            responsible=self.center_b,
            reference="IA-TERR-XLSX-OUT",
            deadline=date.today() - timedelta(days=2),
        )
        self.client.force_login(self.director)
        params = {
            "mode": "current",
            "department": "Departamento A",
            "group_by": "district",
            "quick": "overdue",
        }

        page = self.client.get(reverse("director_statistics"), params)
        export = self.client.get(reverse("director_statistics_xlsx"), params)

        self.assertContains(page, reverse("director_statistics_xlsx"))
        self.assertEqual(export.status_code, 200)
        content = b"".join(export.streaming_content)
        workbook = load_workbook(BytesIO(content), data_only=False)
        center_sheet = workbook["Centros educativos"]
        exported_codes = [
            center_sheet.cell(row=row, column=1).value
            for row in range(5, center_sheet.max_row + 1)
        ]
        self.assertEqual(exported_codes, [self.center_a.code])
        self.assertEqual(workbook["Resumen"]["E5"].value, page.context["center_count"])
        self.assertEqual(
            workbook["Resumen"]["E9"].value,
            page.context["overdue_center_count"],
        )
        log = ActivityLog.objects.get(action="director_statistics_xlsx_exported")
        self.assertEqual(log.details["filters"]["department"], "Departamento A")
        self.assertEqual(log.details["filters"]["quick"], "overdue")
        self.assertEqual(log.details["rows"]["centers"], page.context["center_count"])
