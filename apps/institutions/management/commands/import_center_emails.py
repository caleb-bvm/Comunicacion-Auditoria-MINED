"""Merge the official email directory without activating or notifying centers."""
import hashlib
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.core.validators import validate_email
from django.db import transaction
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.audits.models import ActivityLog
from apps.institutions.models import Organization


class Command(BaseCommand):
    help = "Carga códigos, nombres y correos desde XLSX, conservando los datos y accesos existentes."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_file")
        parser.add_argument("--dry-run", action="store_true")

    def read_rows(self, source):
        try:
            workbook = load_workbook(source, read_only=True, data_only=False)
        except Exception as exc:
            raise CommandError("No se pudo leer el archivo XLSX.") from exc
        try:
            sheet = workbook.active
            iterator = sheet.iter_rows()
            header = [str(cell.value or "").strip() for cell in next(iterator)]
            required = ["Codigo de Centro Escolar", "Nombre de Centro Escolar", "Email Address"]
            if any(name not in header for name in required):
                raise CommandError("El XLSX debe incluir: " + ", ".join(required))
            indices = [header.index(name) for name in required]
            rows, codes, emails = [], set(), set()
            for line, cells in enumerate(iterator, start=2):
                if not any(cell.value is not None for cell in cells):
                    continue
                selected = [cells[i] for i in indices]
                if any(cell.data_type == "f" for cell in selected):
                    raise CommandError(f"Fila {line}: los datos deben ser valores, no fórmulas.")
                code, name, email = [str(cell.value or "").strip() for cell in selected]
                email = email.lower()
                # Codes must remain text: do not lose leading zeroes or guess numeric formatting.
                if not isinstance(selected[0].value, str) or not code or len(code) > 30:
                    raise CommandError(f"Fila {line}: el código debe ser texto de hasta 30 caracteres.")
                try:
                    User._meta.get_field("username").run_validators(code)
                    validate_email(email)
                except ValidationError as exc:
                    raise CommandError(f"Fila {line}: código o correo inválido.") from exc
                if not name or len(name) > 255 or len(email) > 254:
                    raise CommandError(f"Fila {line}: nombre o correo vacío o demasiado largo.")
                if code in codes or email in emails:
                    raise CommandError(f"Fila {line}: código o correo duplicado.")
                codes.add(code)
                emails.add(email)
                rows.append({"code": code, "name": name, "email": email})
            if not rows:
                raise CommandError("El archivo no contiene centros.")
            return rows
        finally:
            workbook.close()

    @transaction.atomic
    def handle(self, *args, **options):
        source = Path(options["xlsx_file"])
        if not source.is_file():
            raise CommandError("No se encontró el archivo XLSX.")
        rows = self.read_rows(source)
        existing = {item.code: item for item in Organization.objects.select_for_update()}
        email_owners = {}
        for item in existing.values():
            if item.email:
                email_owners.setdefault(item.email.lower(), set()).add(item.code)
        accounts = list(User.objects.select_for_update().select_related("organization"))
        by_username = {user.username: user for user in accounts}
        by_center = {}
        for account in accounts:
            if account.role == User.Role.INSTITUTION and account.organization_id:
                by_center.setdefault(account.organization.code, []).append(account)
        counts = {"centros_nuevos": 0, "correos_actualizados": 0, "sin_cambios": 0,
                  "perfiles_actualizados": 0, "usuarios_normalizados": 0,
                  "usuarios_personalizados_conservados": 0}
        plans = []
        for row in rows:
            code, email = row["code"], row["email"]
            center = existing.get(code)
            if center and center.kind != Organization.Kind.EDUCATIONAL_CENTER:
                raise CommandError(f"El código {code} pertenece a una dependencia que no es un centro.")
            if email_owners.get(email, set()) - {code}:
                raise CommandError(f"El correo del código {code} está asignado a otra institución.")
            if any(user.email.lower() == email and (
                user.role != User.Role.INSTITUTION or not user.organization_id
                or user.organization.code != code
            ) for user in accounts):
                raise CommandError(f"El correo del código {code} pertenece a otra cuenta.")
            if center is None:
                counts["centros_nuevos"] += 1
            elif center.email != email:
                counts["correos_actualizados"] += 1
            else:
                counts["sin_cambios"] += 1
            changes = []
            for user in by_center.get(code, []):
                fields = []
                if user.email != email:
                    fields.append("email")
                    counts["perfiles_actualizados"] += 1
                # Preserve custom test logins; normalize only the previous generated convention.
                if user.username == f"centro.{code}":
                    owner = by_username.get(code)
                    if owner and owner.pk != user.pk:
                        raise CommandError(f"El usuario {code} ya existe. Resuelva el conflicto antes de importar.")
                    fields.append("username")
                    counts["usuarios_normalizados"] += 1
                elif user.username != code:
                    counts["usuarios_personalizados_conservados"] += 1
                changes.append((user, fields))
            plans.append((row, center, changes))
        if not options["dry_run"]:
            for row, center, changes in plans:
                if center is None:
                    Organization.objects.create(**row, kind=Organization.Kind.EDUCATIONAL_CENTER)
                elif center.email != row["email"]:
                    center.email = row["email"]
                    center.save(update_fields=["email", "updated_at"])
                for user, fields in changes:
                    if fields:
                        previous_email = user.email
                        user.email = row["email"]
                        if "username" in fields:
                            user.username = row["code"]
                        if previous_email != user.email and user.activation_requested_at:
                            user.activation_requested_at = None
                            fields.append("activation_requested_at")
                        user.save(update_fields=fields)
            if any(counts[key] for key in (
                "centros_nuevos", "correos_actualizados", "perfiles_actualizados", "usuarios_normalizados"
            )):
                ActivityLog.objects.create(
                    action="center_email_directory_imported", target_type="Organization",
                    details={"source_filename": source.name,
                             "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                             "rows": len(rows), **counts},
                )
        mode = "SIMULACIÓN" if options["dry_run"] else "IMPORTACIÓN"
        self.stdout.write(f"{mode}: {len(rows)} centros. " + "; ".join(f"{k}={v}" for k, v in counts.items()))
        self.stdout.write("No se crearon ni activaron cuentas y no se enviaron correos.")
