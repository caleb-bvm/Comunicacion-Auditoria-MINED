from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.institutions.models import Organization


REQUIRED_HEADERS = {
    "CODIGO CE",
    "NOMBRE CE",
    "DEPARTAMENTO",
    "DISTRITO",
    "DIRECCION",
    "LATITUD",
    "LONGITUD",
}
MIN_LATITUDE = Decimal("13")
MAX_LATITUDE = Decimal("15")
MIN_LONGITUDE = Decimal("-91")
MAX_LONGITUDE = Decimal("-87")


class Command(BaseCommand):
    help = "Actualiza ubicación y coordenadas de centros desde el XLSX georreferenciado."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_file", help="Ruta del archivo XLSX georreferenciado.")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Valida y muestra el resultado esperado sin guardar cambios.",
        )

    def handle(self, *args, **options):
        source = Path(options["xlsx_file"])
        if not source.is_file():
            raise CommandError(f"No se encontró el archivo: {source}")

        rows = self.read_rows(source)
        organizations = Organization.objects.in_bulk(field_name="code")
        matched = []
        unmatched = []
        alias_matches = 0
        without_coordinates = 0

        for row in rows:
            organization = organizations.get(row["code"])
            if organization is None:
                alias = organizations.get(f"cod{row['code']}")
                if alias and Organization._normalized_location(alias.name) == row["normalized_name"]:
                    organization = alias
                    alias_matches += 1
            if organization is None:
                unmatched.append(row["code"])
                continue
            if row["latitude"] is None:
                without_coordinates += 1
            matched.append((organization, row))

        updated = 0
        unchanged = 0
        now = timezone.now()
        fields = ["department", "district", "address", "latitude", "longitude", "updated_at"]
        pending = []
        with transaction.atomic():
            locked = {
                item.pk: item
                for item in Organization.objects.select_for_update().filter(
                    pk__in=[organization.pk for organization, _row in matched]
                )
            }
            for organization, row in matched:
                current = locked[organization.pk]
                changes = {}
                for field in ("department", "district", "address"):
                    if row[field] and getattr(current, field) != row[field]:
                        changes[field] = row[field]
                if row["latitude"] is not None:
                    if current.latitude != row["latitude"]:
                        changes["latitude"] = row["latitude"]
                    if current.longitude != row["longitude"]:
                        changes["longitude"] = row["longitude"]
                if not changes:
                    unchanged += 1
                    continue
                updated += 1
                for field, value in changes.items():
                    setattr(current, field, value)
                current.updated_at = now
                pending.append(current)

            if pending and not options["dry_run"]:
                Organization.objects.bulk_update(pending, fields, batch_size=500)

        mode = "SIMULACIÓN" if options["dry_run"] else "IMPORTACIÓN"
        self.stdout.write(
            self.style.SUCCESS(
                f"{mode} completada: filas={len(rows)}, coincidencias={len(matched)}, "
                f"alias_cod={alias_matches}, actualizadas={updated}, sin_cambios={unchanged}, "
                f"sin_coordenadas={without_coordinates}, no_encontradas={len(unmatched)}."
            )
        )
        if unmatched:
            self.stdout.write("Códigos no encontrados: " + ", ".join(unmatched))

    def read_rows(self, source):
        try:
            workbook = load_workbook(source, read_only=True, data_only=False)
        except (OSError, ValueError) as exc:
            raise CommandError(f"No fue posible abrir el XLSX: {exc}") from exc

        try:
            sheet = self.find_sheet(workbook)
            cells = sheet.iter_rows(values_only=False)
            header_cells = next(cells, None)
            if not header_cells:
                raise CommandError("El XLSX está vacío.")
            headers = [str(cell.value or "").strip() for cell in header_cells]
            missing = REQUIRED_HEADERS - set(headers)
            if missing:
                raise CommandError("Faltan columnas obligatorias: " + ", ".join(sorted(missing)))
            positions = {header: headers.index(header) for header in REQUIRED_HEADERS}

            rows = []
            seen = set()
            for line_number, cells in enumerate(cells, start=2):
                code_cell = cells[positions["CODIGO CE"]]
                if code_cell.value in (None, ""):
                    continue
                self.reject_formula(code_cell, line_number, "CODIGO CE")
                code = self.normalize_code(code_cell.value, line_number)
                if code in seen:
                    raise CommandError(f"Fila {line_number}: el código {code} está repetido.")
                seen.add(code)

                name_cell = cells[positions["NOMBRE CE"]]
                self.reject_formula(name_cell, line_number, "NOMBRE CE")
                name = self.clean_text(name_cell.value)
                if not name:
                    raise CommandError(f"Fila {line_number}: NOMBRE CE es obligatorio.")

                latitude = self.coordinate(cells[positions["LATITUD"]], line_number, "LATITUD")
                longitude = self.coordinate(cells[positions["LONGITUD"]], line_number, "LONGITUD")
                if (latitude is None) != (longitude is None):
                    raise CommandError(
                        f"Fila {line_number}: latitud y longitud deben venir juntas."
                    )
                if latitude is not None and not (MIN_LATITUDE <= latitude <= MAX_LATITUDE):
                    raise CommandError(f"Fila {line_number}: latitud fuera de El Salvador.")
                if longitude is not None and not (MIN_LONGITUDE <= longitude <= MAX_LONGITUDE):
                    raise CommandError(f"Fila {line_number}: longitud fuera de El Salvador.")

                rows.append(
                    {
                        "code": code,
                        "normalized_name": Organization._normalized_location(name),
                        "department": self.text_cell(cells[positions["DEPARTAMENTO"]], line_number, "DEPARTAMENTO"),
                        "district": self.text_cell(cells[positions["DISTRITO"]], line_number, "DISTRITO"),
                        "address": self.text_cell(cells[positions["DIRECCION"]], line_number, "DIRECCION"),
                        "latitude": latitude,
                        "longitude": longitude,
                    }
                )
            return rows
        finally:
            workbook.close()

    def find_sheet(self, workbook):
        for sheet in workbook.worksheets:
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = {str(value or "").strip() for value in first_row}
            if REQUIRED_HEADERS.issubset(headers):
                return sheet
        raise CommandError("Ninguna hoja contiene las columnas de georreferencia requeridas.")

    def coordinate(self, cell, line_number, label):
        self.reject_formula(cell, line_number, label)
        if cell.value in (None, ""):
            return None
        try:
            return Decimal(str(cell.value)).quantize(Decimal("0.0000001"))
        except (InvalidOperation, ValueError) as exc:
            raise CommandError(f"Fila {line_number}: {label} debe ser numérica.") from exc

    def text_cell(self, cell, line_number, label):
        self.reject_formula(cell, line_number, label)
        return self.clean_text(cell.value)

    @staticmethod
    def clean_text(value):
        return " ".join(str(value or "").split())

    @staticmethod
    def reject_formula(cell, line_number, label):
        if cell.data_type == "f":
            raise CommandError(f"Fila {line_number}: {label} no puede contener una fórmula.")

    @staticmethod
    def normalize_code(value, line_number):
        if isinstance(value, bool):
            raise CommandError(f"Fila {line_number}: CODIGO CE inválido.")
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        code = str(value).strip()
        if not code or not code.isdigit():
            raise CommandError(f"Fila {line_number}: CODIGO CE debe ser numérico.")
        return code
